#!/usr/bin/env python3
"""The PPTX import oracle (gslides-parity SPEC-5 5.4; R04 10 test 5): python-pptx's own reading of
a file as JSON, so `packages/import/src/pptx/oracle.test.ts` can compare the reader's counts with
an independent library's. Runs only when the fonts venv exists (.turboslide/venv,
scripts/requirements.txt); the test skips with a notice otherwise and the check chain never
installs it.

    .turboslide/venv/bin/python scripts/pptx-oracle.py <file.pptx>

Prints one JSON object: the slide size in EMU, and per slide (in `sldIdLst` order) the shape count
(groups counted by their leaf members, hidden shapes included, the exporter's own chrome and page
raster left out by the same name rule the reader applies), the merge origins of every table, the
chart categories and values, the picture count and the media count.
"""

from __future__ import annotations

import json
import re
import sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn
except ImportError as error:  # pragma: no cover - the venv is opt in
    sys.stderr.write(f"pptx-oracle: {error}. Install the fonts venv first.\n")
    sys.exit(2)

# ts:<slide>#<block>[/part|:n][@...]: the exporter's grammar; the chrome the reader leaves out
NAME_RE = re.compile(r"^ts:([a-z0-9-]+)#([a-z0-9-]+)([/:][^@]*)?((?:@[^@]+)*)$")


def is_exporter_chrome(name: str) -> bool:
    match = NAME_RE.match(name.strip())
    if match is None:
        return False
    block = match.group(2)
    part = (match.group(3) or "")[1:]
    if block in ("counter", "wordmark", "sheet"):
        return True
    if block == "mark" and re.fullmatch(r"\d+", part or ""):
        return True
    if block in ("frame", "cross", "chip") and re.match(r"\d", part or ""):
        return True
    return False


def leaf_shapes(shapes) -> list:
    out = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out.extend(leaf_shapes(shape.shapes))
        else:
            out.append(shape)
    return out


def table_merges(shape) -> int:
    """The merge origin cells of a table (python-pptx: `is_merge_origin`)."""
    count = 0
    for row in shape.table.rows:
        for cell in row.cells:
            if cell.is_merge_origin:
                count += 1
    return count


C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def cache_values(holder):
    """The cached points of a c:cat or c:val holder, by index, as strings."""
    if holder is None:
        return []
    cache = holder.find(f".//{{{C_NS}}}strCache")
    if cache is None:
        cache = holder.find(f".//{{{C_NS}}}numCache")
    if cache is None:
        return []
    points = {}
    for pt in cache.findall(f"{{{C_NS}}}pt"):
        v = pt.find(f"{{{C_NS}}}v")
        points[int(pt.get("idx", "0"))] = "" if v is None or v.text is None else v.text
    count = cache.find(f"{{{C_NS}}}ptCount")
    n = int(count.get("val", "0")) if count is not None else (max(points) + 1 if points else 0)
    return [points.get(i, "") for i in range(max(n, (max(points) + 1) if points else 0))]


def chart_facts(shape) -> dict:
    """The first plot's categories and series from the caches, read from the part's XML (python-pptx
    cannot iterate a 3D plot, and the reader never opens the workbook either)."""
    plot_area = shape.chart._chartSpace.chart.plotArea
    plots = [el for el in plot_area if el.tag.endswith("Chart") and el.tag.startswith(f"{{{C_NS}}}")]
    first = plots[0] if plots else None
    categories = []
    series = []
    if first is not None:
        sers = sorted(first.findall(f"{{{C_NS}}}ser"), key=lambda ser: int(ser.find(f"{{{C_NS}}}idx").get("val", "0")) if ser.find(f"{{{C_NS}}}idx") is not None else 0)
        for index, ser in enumerate(sers):
            if index == 0:
                categories = cache_values(ser.find(f"{{{C_NS}}}cat"))
            values = cache_values(ser.find(f"{{{C_NS}}}val"))
            tx = ser.find(f"{{{C_NS}}}tx")
            name_values = cache_values(tx)
            name = name_values[0] if name_values else f"Series {index + 1}"
            series.append({"name": name, "values": [None if v == "" else float(v) for v in values]})
    return {"kind": first.tag.split("}")[-1] if first is not None else "", "categories": categories, "series": series, "plots": len(plots)}


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        sys.stderr.write("usage: pptx-oracle.py <file.pptx>\n")
        return 2
    prs = Presentation(argv[0])
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        leaves = leaf_shapes(slide.shapes)
        content = [s for s in leaves if not is_exporter_chrome(s.name)]
        # a mc:AlternateContent wrapper is not a shape to python-pptx: count its Fallback objects
        alternates = slide.shapes._spTree.findall("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
        fallback_shapes = 0
        for alt in alternates:
            fallback = alt.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback")
            if fallback is not None:
                fallback_shapes += len([c for c in fallback if c.tag in (qn("p:sp"), qn("p:pic"), qn("p:graphicFrame"), qn("p:cxnSp"))])
        tables = [s for s in content if getattr(s, "has_table", False)]
        charts = [s for s in content if getattr(s, "has_chart", False)]
        # a media frame is a p:pic whose p:nvPr carries a:videoFile or a:audioFile (python-pptx reads a retagged audio as a picture)
        def is_media(shape) -> bool:
            nvPr = shape._element.find(".//" + qn("p:nvPr"))
            return nvPr is not None and (nvPr.find(qn("a:videoFile")) is not None or nvPr.find(qn("a:audioFile")) is not None)

        media = [s for s in content if is_media(s)]
        pictures = [s for s in content if s.shape_type == MSO_SHAPE_TYPE.PICTURE and not is_media(s)]
        slides.append(
            {
                "index": index,
                "id": int(slide.slide_id),
                "shapes": len(content) + fallback_shapes,
                "hidden": len([s for s in content if s._element.xpath(".//p:cNvPr")[0].get("hidden") == "1"]),
                "tables": [{"rows": len(s.table.rows), "columns": len(s.table.columns), "merges": table_merges(s)} for s in tables],
                "charts": [chart_facts(s) for s in charts],
                "pictures": len(pictures),
                "media": len(media),
            }
        )
    facts = {"size": {"cx": int(prs.slide_width), "cy": int(prs.slide_height)}, "slides": slides}
    sys.stdout.write(json.dumps(facts, indent=2, sort_keys=True) + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
